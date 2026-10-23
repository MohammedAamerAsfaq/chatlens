from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import uuid4

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from rest_framework.test import APIClient
from openpyxl import Workbook

from apps.real_estate.models import (
    ActivityType, ClassificationStatus, Listing, OpportunityItem, Party, Property,
    PropertyCandidate, PropertyRequirement, PropertyType,
    RealEstateMockSourceRecord, RealEstateOpportunity, RealEstatePromptConfig,
)
from apps.task_management.models import BackgroundTask
from apps.tenancy.models import Company, CompanyMembership
from apps.real_estate.services.property_types import seed_default_property_types
from apps.real_estate.services.extraction_projection import project_extraction


class RealEstateApiTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name='API Estate', slug='api-estate', industry_type='real_estate')
        self.other = Company.objects.create(name='Other Estate', slug='other-estate', industry_type='real_estate')
        self.types = {item.code: item for item in seed_default_property_types(self.company)}
        self.other_types = {item.code: item for item in seed_default_property_types(self.other)}
        self.user = get_user_model().objects.create_user('api-estate-owner')
        CompanyMembership.objects.create(
            company=self.company, user=self.user, role=CompanyMembership.ROLE_SUPER_USER,
        )
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def test_real_estate_prompt_api_lists_defaults_and_saves_company_override(self):
        response = self.client.get('/api/real-estate/prompts/')
        self.assertEqual(response.status_code, 200, response.json())
        self.assertEqual(
            {row['key'] for row in response.json()},
            {'gate', 'extraction', 'match_rerank'},
        )
        self.assertTrue(all(row['is_default'] for row in response.json()))

        response = self.client.patch('/api/real-estate/prompts/extraction/', {
            'body': 'Company-specific Real Estate extraction instruction.',
        }, format='json')
        self.assertEqual(response.status_code, 200, response.json())
        self.assertFalse(response.json()['is_default'])
        prompt = RealEstatePromptConfig.objects.get(company=self.company, key='extraction')
        self.assertEqual(prompt.body, 'Company-specific Real Estate extraction instruction.')
        self.assertFalse(RealEstatePromptConfig.objects.filter(company=self.other).exists())

    def test_card_field_config_defaults_empty_and_saves_company_scoped_list(self):
        response = self.client.get('/api/real-estate/prompts/card-field-config/')
        self.assertEqual(response.status_code, 200, response.json())
        self.assertEqual(response.json(), {'fields': []})

        response = self.client.put(
            '/api/real-estate/prompts/card-field-config/',
            {'fields': ['party_name', 'property_type', 'location']}, format='json',
        )
        self.assertEqual(response.status_code, 200, response.json())
        self.assertEqual(response.json(), {'fields': ['party_name', 'property_type', 'location']})

        from apps.real_estate.models import RealEstateSettings
        self.assertEqual(
            RealEstateSettings.objects.get(company=self.company).card_field_config,
            ['party_name', 'property_type', 'location'],
        )

    def test_card_field_config_rejects_non_string_list(self):
        response = self.client.put(
            '/api/real-estate/prompts/card-field-config/', {'fields': ['ok', 5]}, format='json',
        )
        self.assertEqual(response.status_code, 400)

    def test_card_action_config_defaults_and_saves(self):
        response = self.client.get('/api/real-estate/prompts/card-action-config/')
        self.assertEqual(response.status_code, 200, response.json())
        self.assertEqual(response.json(), {'statuses': [], 'max_inline': 2})

        response = self.client.put(
            '/api/real-estate/prompts/card-action-config/',
            {'statuses': ['qualified', 'won', 'lost'], 'max_inline': 2}, format='json',
        )
        self.assertEqual(response.status_code, 200, response.json())
        self.assertEqual(response.json(), {'statuses': ['qualified', 'won', 'lost'], 'max_inline': 2})

    def test_card_action_config_rejects_unknown_status(self):
        response = self.client.put(
            '/api/real-estate/prompts/card-action-config/',
            {'statuses': ['not_a_real_status'], 'max_inline': 2}, format='json',
        )
        self.assertEqual(response.status_code, 400)

    def test_card_configs_readable_by_any_member_but_writable_only_by_admins(self):
        member_user = get_user_model().objects.create_user('api-estate-member')
        CompanyMembership.objects.create(
            company=self.company, user=member_user, role=CompanyMembership.ROLE_USER,
        )
        member_client = APIClient()
        member_client.force_authenticate(member_user)

        response = member_client.get('/api/real-estate/prompts/card-field-config/')
        self.assertEqual(response.status_code, 200, response.json())
        response = member_client.get('/api/real-estate/prompts/card-action-config/')
        self.assertEqual(response.status_code, 200, response.json())

        response = member_client.put(
            '/api/real-estate/prompts/card-field-config/', {'fields': ['party_name']}, format='json',
        )
        self.assertEqual(response.status_code, 403)
        response = member_client.put(
            '/api/real-estate/prompts/card-action-config/', {'statuses': [], 'max_inline': 2}, format='json',
        )
        self.assertEqual(response.status_code, 403)

    def test_card_configs_are_company_scoped(self):
        other_user = get_user_model().objects.create_user('api-estate-other-owner')
        CompanyMembership.objects.create(
            company=self.other, user=other_user, role=CompanyMembership.ROLE_SUPER_USER,
        )
        self.client.put(
            '/api/real-estate/prompts/card-field-config/', {'fields': ['party_name']}, format='json',
        )
        other_client = APIClient()
        other_client.force_authenticate(other_user)
        response = other_client.get('/api/real-estate/prompts/card-field-config/')
        self.assertEqual(response.json(), {'fields': []})

    def test_properties_are_company_scoped_and_archived_on_delete(self):
        visible = Property.objects.create(company=self.company, title='Visible flat', property_type=self.types['apartment'])
        listing = Listing.objects.create(
            company=self.company, property=visible, listing_type='sale', asking_price=100,
            is_verified=True, status='active',
        )
        Property.objects.create(company=self.other, title='Hidden villa', property_type=self.other_types['villa'])

        response = self.client.get('/api/real-estate/properties/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual([row['id'] for row in response.json()['results']], [visible.pk])

        response = self.client.delete(f'/api/real-estate/properties/{visible.pk}/')
        self.assertEqual(response.status_code, 204)
        visible.refresh_from_db()
        self.assertTrue(visible.is_archived)
        self.assertEqual(visible.availability_status, 'inactive')
        listing.refresh_from_db()
        self.assertTrue(listing.is_archived)
        self.assertEqual(listing.status, 'closed')

    def test_property_create_ignores_client_company_and_uses_active_company(self):
        response = self.client.post('/api/real-estate/properties/', {
            'company': self.other.pk,
            'title': 'Owned flat',
            'property_type': 'apartment',
        }, format='json')
        self.assertEqual(response.status_code, 201, response.json())
        self.assertEqual(Property.objects.get(pk=response.json()['id']).company, self.company)

    def test_property_hierarchy_options_are_registered_and_company_scoped(self):
        visible = Property.objects.create(
            company=self.company, title='Local building', property_type=self.types['building'],
        )
        Property.objects.create(
            company=self.other, title='Foreign building', property_type=self.other_types['building'],
        )

        response = self.client.get('/api/real-estate/properties/hierarchy-options/')

        self.assertEqual(response.status_code, 200, response.json())
        self.assertEqual(response.json(), [{
            'id': visible.pk,
            'title': 'Local building',
            'reference_code': '',
            'parent_id': None,
        }])

    def test_party_endpoints_are_company_scoped_and_reject_foreign_links(self):
        property_record = Property.objects.create(
            company=self.company, title='Scoped villa', property_type=self.types['villa'],
        )
        visible = Party.objects.create(company=self.company, first_name='Visible owner')
        hidden = Party.objects.create(company=self.other, first_name='Hidden owner')

        response = self.client.get('/api/real-estate/parties/')
        self.assertEqual(response.status_code, 200, response.json())
        self.assertEqual([row['id'] for row in response.json()['results']], [visible.pk])

        response = self.client.post('/api/real-estate/property-party-associations/', {
            'property': property_record.pk,
            'party': hidden.pk,
            'relationship_type': 'current_owner',
        }, format='json')
        self.assertEqual(response.status_code, 400, response.json())

    def test_mock_source_preview_returns_model_shaped_rows_without_persisting(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Prospects'
        sheet.append(['Name', 'Mobile', 'Property', 'Budget'])
        sheet.append(['Demo Buyer', '971500000000', 'Dubai Marina apartment', 2500000])
        content = BytesIO()
        workbook.save(content)
        property_count = Property.objects.count()

        response = self.client.post('/api/real-estate/properties/mock-source-preview/', {
            'files': SimpleUploadedFile(
                'prospects.xlsx', content.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            ),
            'record_limit': 1,
        }, format='multipart')

        self.assertEqual(response.status_code, 200, response.json())
        payload = response.json()
        self.assertFalse(payload['persisted'])
        self.assertEqual(payload['count'], 1)
        self.assertEqual(payload['record_limit'], 1)
        self.assertEqual(payload['rows'][0]['source_file_name'], 'prospects.xlsx')
        self.assertEqual(payload['rows'][0]['source_sheet_name'], 'Prospects')
        self.assertEqual(payload['rows'][0]['source_row_number'], 1)
        self.assertEqual(payload['rows'][0]['row_kind'], 'header')
        self.assertIsNone(payload['rows'][0]['imported_at'])
        self.assertEqual(Property.objects.count(), property_count)

    def test_mock_source_preview_rejects_unsupported_record_limit(self):
        workbook = Workbook()
        workbook.active.append(['Name'])
        content = BytesIO()
        workbook.save(content)

        response = self.client.post('/api/real-estate/properties/mock-source-preview/', {
            'files': SimpleUploadedFile('prospects.xlsx', content.getvalue()),
            'record_limit': 2,
        }, format='multipart')

        self.assertEqual(response.status_code, 400)
        self.assertIn('record_limit', response.json())

    def test_mock_source_files_can_be_queued_imported_and_deleted(self):
        from apps.real_estate.services.mock_source_import import import_mock_sources

        with TemporaryDirectory() as temp_dir:
            source_dir = Path(temp_dir) / 'static' / 'Real Estate Data'
            source_dir.mkdir(parents=True)
            workbook = Workbook()
            workbook.active.append(['Name', 'Mobile', 'Property'])
            workbook.active.append(['Demo Buyer', '971500000000', 'Marina flat'])
            workbook.save(source_dir / 'sample.xlsx')

            with override_settings(BASE_DIR=temp_dir):
                response = self.client.get('/api/real-estate/mock-source-records/available-files/')
                self.assertEqual(response.status_code, 200, response.json())
                self.assertEqual([row['name'] for row in response.json()['files']], ['sample.xlsx'])

                response = self.client.post('/api/real-estate/mock-source-records/import-files/', {
                    'import_all': False, 'file_names': ['sample.xlsx'],
                }, format='json')
                self.assertEqual(response.status_code, 202, response.json())
                task = BackgroundTask.objects.get(pk=response.json()['task_id'])
                self.assertEqual(task.queue_name, 'real_estate_imports')
                self.assertEqual(task.company, self.company)

                first = import_mock_sources(self.company, uuid4(), ['sample.xlsx'])
                second = import_mock_sources(self.company, uuid4(), ['sample.xlsx'])
                self.assertEqual(first['records_created'], 2)
                self.assertEqual(second['records_created'], 0)
                self.assertEqual(second['duplicates_skipped'], 2)

                listed = self.client.get('/api/real-estate/mock-source-records/?page_size=10')
                self.assertEqual(listed.status_code, 200, listed.json())
                self.assertEqual(listed.json()['count'], 2)
                self.assertEqual(listed.json()['results'][0]['company_id'], self.company.pk)

                blocked = self.client.delete('/api/real-estate/mock-source-records/delete-all/')
                self.assertEqual(blocked.status_code, 409)
                task.status = BackgroundTask.STATUS_SUCCEEDED
                task.save(update_fields=['status'])
                deleted = self.client.delete('/api/real-estate/mock-source-records/delete-all/')
                self.assertEqual(deleted.status_code, 200, deleted.json())
                self.assertEqual(deleted.json()['deleted'], 2)

    def test_mock_source_delete_all_is_company_scoped(self):
        common = {
            'import_run_id': uuid4(), 'source_file_name': 'source.xlsx',
            'source_file_checksum': 'a' * 64, 'source_sheet_name': 'Sheet1',
            'source_sheet_index': 0, 'source_row_number': 1,
            'raw_values': ['Name'], 'display_values': ['Name'],
            'row_checksum': 'b' * 64,
        }
        RealEstateMockSourceRecord.objects.create(company=self.company, **common)
        RealEstateMockSourceRecord.objects.create(company=self.other, **common)

        response = self.client.delete('/api/real-estate/mock-source-records/delete-all/')

        self.assertEqual(response.status_code, 200, response.json())
        self.assertFalse(RealEstateMockSourceRecord.objects.filter(company=self.company).exists())
        self.assertTrue(RealEstateMockSourceRecord.objects.filter(company=self.other).exists())

    def test_property_inventory_filters_keep_tracking_independent(self):
        own = Property.objects.create(
            company=self.company, title='Own unit', property_type=self.types['apartment'],
            inventory_source_type='own', is_tracked=True,
        )
        authorized = Property.objects.create(
            company=self.company, title='Authorized unit', property_type=self.types['villa'],
            inventory_source_type='authorized', is_tracked=False,
        )
        response = self.client.get('/api/real-estate/properties/?inventory_type=authorized')
        self.assertEqual([row['id'] for row in response.json()['results']], [authorized.pk])
        response = self.client.get('/api/real-estate/properties/?inventory_type=tracked')
        self.assertEqual([row['id'] for row in response.json()['results']], [own.pk])

    def test_property_type_endpoint_exposes_system_schema(self):
        response = self.client.get('/api/real-estate/property-types/?page_size=100')
        self.assertEqual(response.status_code, 200)
        apartment = next(row for row in response.json()['results'] if row['code'] == 'apartment')
        self.assertIn('flat', apartment['aliases'])
        self.assertIn('floor', {field['code'] for field in apartment['field_definitions']})

    def test_property_type_master_supports_custom_types_and_safe_deactivation(self):
        response = self.client.post('/api/real-estate/property-types/', {
            'name': 'Hotel Apartment', 'code': 'hotel_apartment',
            'category': 'residential', 'sort_order': 25,
        }, format='json')
        self.assertEqual(response.status_code, 201, response.json())
        type_id = response.json()['id']
        self.assertEqual(PropertyType.objects.get(pk=type_id).company, self.company)

        response = self.client.post('/api/real-estate/properties/', {
            'title': 'Serviced unit', 'property_type': 'hotel_apartment', 'area_unit': 'acre',
        }, format='json')
        self.assertEqual(response.status_code, 201, response.json())
        self.assertEqual(response.json()['property_type_name'], 'Hotel Apartment')
        self.assertEqual(response.json()['area_unit'], 'acre')

        self.assertEqual(self.client.delete(f'/api/real-estate/property-types/{type_id}/').status_code, 204)
        response = self.client.post('/api/real-estate/properties/', {
            'title': 'Another serviced unit', 'property_type': 'hotel_apartment',
        }, format='json')
        self.assertEqual(response.status_code, 400)
        self.assertTrue(PropertyType.objects.get(pk=type_id).is_active is False)

    def test_candidate_promotion_is_explicit(self):
        opportunity = RealEstateOpportunity.objects.create(
            company=self.company,
            activity_type=ActivityType.SELLING,
            classification_status=ClassificationStatus.CLASSIFIED,
            first_seen_at=timezone.now(),
        )
        candidate = PropertyCandidate.objects.create(
            company=self.company,
            opportunity=opportunity,
            title='Seller villa',
            property_type=self.types['villa'],
            asking_price=3000000,
        )
        self.assertFalse(Property.objects.filter(company=self.company).exists())
        response = self.client.post(
            f'/api/real-estate/candidates/{candidate.pk}/promote/', {}, format='json',
        )
        self.assertEqual(response.status_code, 201, response.json())
        candidate.refresh_from_db()
        self.assertEqual(candidate.review_status, 'promoted')
        self.assertTrue(Property.objects.filter(pk=response.json()['property_id']).exists())

    def test_non_real_estate_company_is_blocked(self):
        self.company.industry_type = Company.INDUSTRY_GENERAL
        self.company.save(update_fields=['industry_type'])
        response = self.client.get('/api/real-estate/properties/')
        self.assertEqual(response.status_code, 403)

    def test_agent_sees_only_assigned_opportunities(self):
        agent = get_user_model().objects.create_user('estate-agent')
        CompanyMembership.objects.create(
            company=self.company, user=agent, role=CompanyMembership.ROLE_USER,
        )
        assigned = RealEstateOpportunity.objects.create(
            company=self.company, assigned_to=agent, activity_type=ActivityType.BUYING,
            classification_status=ClassificationStatus.CLASSIFIED, first_seen_at=timezone.now(),
        )
        RealEstateOpportunity.objects.create(
            company=self.company, assigned_to=self.user, activity_type=ActivityType.SELLING,
            classification_status=ClassificationStatus.CLASSIFIED, first_seen_at=timezone.now(),
        )
        self.client.force_authenticate(agent)
        response = self.client.get('/api/real-estate/opportunities/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual([row['id'] for row in response.json()['results']], [assigned.pk])

    def test_opportunity_board_filters_searches_and_embeds_lane_detail(self):
        buying = RealEstateOpportunity.objects.create(
            company=self.company, activity_type=ActivityType.BUYING,
            classification_status=ClassificationStatus.CLASSIFIED,
            summary='Buyer needs a Marina apartment', original_message='WTB Marina two bedroom',
            first_seen_at=timezone.now(),
        )
        project_extraction(buying, {'properties': [{
            'property_type': 'apartment', 'locations': ['Dubai Marina'],
            'bedrooms': 2, 'budget_max': 2500000,
        }]})
        RealEstateOpportunity.objects.create(
            company=self.company, activity_type=ActivityType.SELLING,
            classification_status=ClassificationStatus.CLASSIFIED,
            summary='Seller offers a villa', first_seen_at=timezone.now(),
        )
        RealEstateOpportunity.objects.create(
            company=self.other, activity_type=ActivityType.BUYING,
            classification_status=ClassificationStatus.CLASSIFIED,
            summary='Other company Marina requirement', first_seen_at=timezone.now(),
        )

        response = self.client.get(
            '/api/real-estate/opportunities/?activity_type=buying&classification_status=classified&search=Marina',
        )

        self.assertEqual(response.status_code, 200, response.json())
        self.assertEqual(response.json()['count'], 1)
        row = response.json()['results'][0]
        self.assertEqual(row['id'], buying.pk)
        self.assertEqual(row['requirement']['acceptable_locations'], ['Dubai Marina'])
        self.assertIsNone(row['property_candidate'])
        self.assertEqual(len(row['items']), 1)
        self.assertEqual(row['items'][0]['location_raw'], 'Dubai Marina')

        response = self.client.patch(
            f'/api/real-estate/opportunities/{buying.pk}/', {'status': 'qualified'}, format='json',
        )
        self.assertEqual(response.status_code, 200, response.json())
        self.assertEqual(response.json()['status'], 'qualified')

    def test_supply_item_can_create_property_without_creating_listing(self):
        opportunity = RealEstateOpportunity.objects.create(
            company=self.company, activity_type=ActivityType.SELLING,
            classification_status=ClassificationStatus.CLASSIFIED,
            summary='Owner offers a Marina apartment', first_seen_at=timezone.now(),
        )
        project_extraction(opportunity, {'properties': [{
            'title': 'Marina apartment', 'property_type': 'apartment',
            'location': 'Dubai Marina', 'asking_price': 2200000,
        }]})
        item = opportunity.items.get()

        response = self.client.post(
            f'/api/real-estate/opportunity-items/{item.pk}/create-property/', {}, format='json',
        )

        self.assertEqual(response.status_code, 201, response.json())
        item.refresh_from_db()
        self.assertEqual(item.match_status, OpportunityItem.MATCH_LINKED)
        self.assertEqual(item.property.title, 'Marina apartment')
        self.assertEqual(item.property.inventory_source_type, 'floating')
        self.assertEqual(item.property.address.normalized_location, 'Dubai Marina')
        self.assertFalse(Listing.objects.filter(property=item.property).exists())

    def test_demand_item_cannot_create_physical_property(self):
        opportunity = RealEstateOpportunity.objects.create(
            company=self.company, activity_type=ActivityType.BUYING,
            classification_status=ClassificationStatus.CLASSIFIED,
            first_seen_at=timezone.now(),
        )
        project_extraction(opportunity, {'properties': [{
            'property_type': 'apartment', 'location': 'Dubai Marina',
        }]})

        response = self.client.post(
            f'/api/real-estate/opportunity-items/{opportunity.items.get().pk}/create-property/',
            {}, format='json',
        )

        self.assertEqual(response.status_code, 400, response.json())
        self.assertFalse(Property.objects.filter(company=self.company).exists())

    def test_property_market_activity_reports_linked_opportunities(self):
        property_record = Property.objects.create(
            company=self.company, title='Linked villa', property_type=self.types['villa'],
        )
        opportunity = RealEstateOpportunity.objects.create(
            company=self.company, activity_type=ActivityType.SELLING,
            classification_status=ClassificationStatus.CLASSIFIED,
            summary='Owner offers a linked villa', first_seen_at=timezone.now(),
        )
        item = OpportunityItem.objects.create(
            company=self.company, opportunity=opportunity, sequence=1,
            item_kind=OpportunityItem.KIND_SUPPLY,
            amount_role=OpportunityItem.AMOUNT_ASKING,
            property=property_record, match_status=OpportunityItem.MATCH_LINKED,
            link_method=OpportunityItem.LINK_MANUAL,
        )

        response = self.client.get(
            f'/api/real-estate/properties/{property_record.pk}/market-activity/',
        )

        self.assertEqual(response.status_code, 200, response.json())
        self.assertEqual(response.json()['total'], 1)
        self.assertEqual(response.json()['counts']['selling'], 1)
        self.assertEqual(response.json()['recent'][0]['item_id'], item.pk)

    def test_csv_preview_then_import_is_explicit(self):
        csv_file = SimpleUploadedFile(
            'properties.csv',
            b'reference_code,title,property_type,city,listing_type,asking_price,is_verified\n'
            b'SV-001,Marina Flat,apartment,Dubai,sale,2200000,true\n',
            content_type='text/csv',
        )
        response = self.client.post(
            '/api/real-estate/properties/import-preview/', {'file': csv_file}, format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.json())
        self.assertFalse(Property.objects.filter(company=self.company).exists())
        row = response.json()['rows'][0]
        self.assertTrue(row['valid'])

        response = self.client.post(
            '/api/real-estate/properties/bulk-import/', {'rows': [row]}, format='json',
        )
        self.assertEqual(response.status_code, 201, response.json())
        record = Property.objects.get(company=self.company, reference_code='SV-001')
        self.assertEqual(record.address.city, 'Dubai')
        self.assertEqual(record.listings.get().asking_price, 2200000)

    def test_csv_alias_maps_to_canonical_type_and_preserves_source_label(self):
        upload = SimpleUploadedFile(
            'properties.csv',
            b'reference_code,title,property_type,inventory_source_type,is_tracked\n'
            b'SV-FLAT,Imported Flat,flat,authorized,true\n',
            content_type='text/csv',
        )
        preview = self.client.post(
            '/api/real-estate/properties/import-preview/', {'file': upload}, format='multipart',
        ).json()['rows'][0]
        self.assertEqual(preview['property_type'], 'apartment')
        self.assertEqual(preview['source_property_type_label'], 'flat')

        response = self.client.post(
            '/api/real-estate/properties/bulk-import/', {'rows': [preview]}, format='json',
        )
        self.assertEqual(response.status_code, 201, response.json())
        record = Property.objects.get(reference_code='SV-FLAT')
        self.assertEqual(record.property_type.code, 'apartment')
        self.assertEqual(record.source_property_type_label, 'flat')
        self.assertEqual(record.inventory_source_type, 'authorized')
        self.assertTrue(record.is_tracked)

    def test_csv_import_skips_exact_duplicate_and_export_is_company_scoped(self):
        existing = Property.objects.create(
            company=self.company, reference_code='SV-002', title='Existing Villa', property_type=self.types['villa'],
        )
        Property.objects.create(company=self.other, reference_code='HIDDEN', title='Hidden', property_type=self.other_types['villa'])
        response = self.client.post('/api/real-estate/properties/bulk-import/', {
            'rows': [{'row_number': 2, 'reference_code': 'SV-002', 'title': 'Existing Villa', 'property_type': 'villa'}],
        }, format='json')
        self.assertEqual(response.status_code, 201, response.json())
        self.assertEqual(response.json()['skipped_rows'], [2])
        self.assertEqual(Property.objects.filter(company=self.company, reference_code='SV-002').count(), 1)

        response = self.client.get('/api/real-estate/properties/export/')
        self.assertEqual(response.status_code, 200)
        body = b''.join(response.streaming_content).decode()
        self.assertIn(existing.title, body)
        self.assertNotIn('Hidden', body)

    def test_csv_preview_rejects_duplicates_within_the_upload(self):
        upload = SimpleUploadedFile(
            'properties.csv',
            b'reference_code,title,property_type\n'
            b'REF-DUP,Marina View,apartment\n'
            b'REF-DUP,Another Title,villa\n',
            content_type='text/csv',
        )

        response = self.client.post(
            '/api/real-estate/properties/import-preview/', {'file': upload}, format='multipart',
        )

        self.assertEqual(response.status_code, 200, response.json())
        self.assertTrue(response.json()['rows'][0]['valid'])
        self.assertFalse(response.json()['rows'][1]['valid'])
        self.assertIn('duplicate', response.json()['rows'][1]['errors'])
