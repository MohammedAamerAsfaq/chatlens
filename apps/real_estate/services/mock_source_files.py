import hashlib
import json
from datetime import date, datetime, time
from pathlib import Path

from django.conf import settings
from rest_framework.exceptions import ValidationError


HEADER_TERMS = {
    'address', 'area', 'budget', 'building', 'customer', 'email', 'location',
    'mobile', 'name', 'owner', 'phone', 'price', 'project', 'property',
    'service', 'status', 'type', 'unit',
}


def mock_source_root():
    return Path(settings.BASE_DIR) / 'static' / 'Real Estate Data'


def available_workbooks():
    root = mock_source_root()
    if not root.exists():
        return []
    return [
        {
            'name': path.relative_to(root).as_posix(),
            'size': path.stat().st_size,
            'modified_at': datetime.fromtimestamp(path.stat().st_mtime).isoformat(),
        }
        for path in sorted(root.rglob('*.xlsx'))
        if not path.name.startswith('~$')
    ]


def resolve_workbooks(file_names=None):
    available = {item['name']: item for item in available_workbooks()}
    selected = list(available) if file_names is None else file_names
    if not selected:
        raise ValidationError({'file_names': 'Select at least one available workbook.'})
    if any(not isinstance(name, str) for name in selected):
        raise ValidationError({'file_names': 'Workbook names must be strings.'})
    selected = list(dict.fromkeys(selected))
    invalid = [name for name in selected if name not in available]
    if invalid:
        raise ValidationError({'file_names': f'Unknown workbook: {invalid[0]}'})
    root = mock_source_root()
    return [(name, root / Path(name)) for name in selected]


def file_checksum(source):
    digest = hashlib.sha256()
    if isinstance(source, Path):
        with source.open('rb') as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b''):
                digest.update(chunk)
        return digest.hexdigest()
    source.seek(0)
    for chunk in iter(lambda: source.read(1024 * 1024), b''):
        digest.update(chunk)
    source.seek(0)
    return digest.hexdigest()


def iter_workbook_rows(company_id, run_id, source, file_name, checksum=None):
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise ValidationError({
            'files': 'Excel import is unavailable because openpyxl is not installed.',
        }) from exc
    checksum = checksum or file_checksum(source)
    try:
        workbook = load_workbook(source, read_only=True, data_only=False)
    except Exception as exc:
        raise ValidationError({'files': f'{file_name}: unable to read workbook ({exc}).'}) from exc
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets):
            first_non_empty = True
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                raw_values = [_json_value(value) for value in values]
                if not any(value not in (None, '') for value in raw_values):
                    continue
                yield {
                    'company_id': company_id,
                    'import_run_id': run_id,
                    'source_file_name': file_name,
                    'source_file_checksum': checksum,
                    'source_sheet_name': sheet.title,
                    'source_sheet_index': sheet_index,
                    'source_row_number': row_number,
                    'row_kind': row_kind(sheet.title, raw_values, first_non_empty),
                    'raw_values': raw_values,
                    'display_values': [_display_value(value) for value in values],
                    'row_checksum': row_checksum(raw_values),
                    'processing_status': 'pending',
                    'processing_notes': '',
                }
                first_non_empty = False
    finally:
        workbook.close()


def row_kind(sheet_name, values, first_non_empty):
    if 'summary' in sheet_name.lower():
        return 'summary'
    if first_non_empty:
        words = {
            word.lower().strip(' _-')
            for value in values if isinstance(value, str)
            for word in value.replace('/', ' ').split()
        }
        if words & HEADER_TERMS:
            return 'header'
    return 'data'


def _json_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def _display_value(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.isoformat(sep=' ')
    if isinstance(value, (date, time)):
        return value.isoformat()
    return str(value)


def row_checksum(values):
    payload = json.dumps(values, ensure_ascii=False, separators=(',', ':'), sort_keys=True)
    return hashlib.sha256(payload.encode('utf-8')).hexdigest()
